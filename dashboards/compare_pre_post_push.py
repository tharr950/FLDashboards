#!/usr/bin/env python3
"""
compare_pre_post_push.py -- one-off diagnostic, NOT part of the regular
pipeline.

The previous scan showed ~2387 MonthlyMetric rows with Tier/Faculty
Leader/Adjunct stored as formula text (e.g. "=VLOOKUP(...)"). That
alone is NOT proof our push broke anything -- openpyxl's default read
mode (data_only=False) always shows the formula string for a formula
cell, whether or not a cached calculated value exists underneath it.
Some of those rows go back to May 2025, long before today's push, so
formula-text-by-itself is expected/normal and not new damage.

The real test: does the CACHED value (data_only=True) for those
formula cells still exist? And did our push change that? This script:

  1. Finds the git commit history for December_Annual_Reviews.xlsx in
     GITHUB_CODE_REPO.
  2. Identifies today's push commit (message starts with
     "data: append MonthlyMetric rows") and the commit immediately
     before it (the pre-push state).
  3. For a handful of sample (Tutor Name, Date Range) rows that show
     formula text for Tier, loads BOTH versions with data_only=True
     and prints the cached Tier/Faculty Leader/Adjunct value in each.

If pre-push has a real cached value (e.g. "Distinguished") and
post-push shows None for the same cell, that confirms our push
stripped the cache. If pre-push ALSO shows None, the cache was already
gone before we touched it and our push isn't the cause.

Usage:
    python3 compare_pre_post_push.py
"""

import base64
import io
import os

import openpyxl
import requests
from dotenv import load_dotenv

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(SCRIPT_DIR, ".env"))

GITHUB_TOKEN = os.environ["GITHUB_TOKEN"]
GITHUB_CODE_REPO = os.environ["GITHUB_CODE_REPO"]
HEADERS = {"Authorization": f"token {GITHUB_TOKEN}", "Accept": "application/vnd.github.v3+json"}

FILE_PATH = "December_Annual_Reviews.xlsx"
SHEET_NAME = "MonthlyMetric"

# Sample rows to check -- picked because the earlier scan showed these
# periods/tutors with formula-text Tier values.
SAMPLE_KEYS = [
    ("Aaron Schneberger", "6/14/26 - 7/11/26"),
    ("Akeema Williams", "6/14/26 - 7/11/26"),
    ("Aaron Schneberger", "5/17/26 - 6/13/26"),
    ("Aaron Schneberger", "12/28/25 - 1/25/26"),
]


def get_commit_history():
    url = f"https://api.github.com/repos/{GITHUB_CODE_REPO}/commits"
    r = requests.get(url, headers=HEADERS, params={"path": FILE_PATH, "per_page": 20}, timeout=30)
    r.raise_for_status()
    return r.json()


def fetch_workbook_at_ref(ref):
    url = f"https://api.github.com/repos/{GITHUB_CODE_REPO}/contents/{FILE_PATH}"
    r = requests.get(url, headers=HEADERS, params={"ref": ref}, timeout=30)
    r.raise_for_status()
    content = base64.b64decode(r.json()["content"])
    return openpyxl.load_workbook(io.BytesIO(content), data_only=True)


def lookup_rows(wb, keys):
    ws = wb[SHEET_NAME]
    header = [c.value for c in ws[1]]
    tutor_col = header.index("Tutor Name") + 1
    date_col = header.index("Date Range") + 1
    tier_col = header.index("Tier") + 1
    fl_col = header.index("Faculty Leader") + 1
    adj_col = header.index("Adjunct/Professional") + 1

    results = {}
    for row in range(2, ws.max_row + 1):
        t = ws.cell(row=row, column=tutor_col).value
        d = ws.cell(row=row, column=date_col).value
        if (t, d) in keys:
            results[(t, d)] = {
                "Tier": ws.cell(row=row, column=tier_col).value,
                "Faculty Leader": ws.cell(row=row, column=fl_col).value,
                "Adjunct/Professional": ws.cell(row=row, column=adj_col).value,
            }
    return results


def main():
    commits = get_commit_history()
    print("Recent commits touching this file:")
    for c in commits[:10]:
        print(f"  {c['sha'][:10]}  {c['commit']['author']['date']}  {c['commit']['message'].splitlines()[0]}")
    print()

    push_idx = None
    for i, c in enumerate(commits):
        if c["commit"]["message"].startswith("data: append MonthlyMetric rows"):
            push_idx = i
            break

    if push_idx is None:
        print("Could not find today's push commit by message prefix -- aborting.")
        return

    post_sha = commits[push_idx]["sha"]
    if push_idx + 1 >= len(commits):
        print("No earlier commit found before the push -- aborting.")
        return
    pre_sha = commits[push_idx + 1]["sha"]

    print(f"Post-push commit: {post_sha[:10]} ({commits[push_idx]['commit']['author']['date']})")
    print(f"Pre-push commit:  {pre_sha[:10]} ({commits[push_idx + 1]['commit']['author']['date']})")
    print()

    print("Loading pre-push workbook (data_only=True)...")
    wb_pre = fetch_workbook_at_ref(pre_sha)
    print("Loading post-push workbook (data_only=True)...")
    wb_post = fetch_workbook_at_ref(post_sha)

    keys = set(SAMPLE_KEYS)
    pre_vals = lookup_rows(wb_pre, keys)
    post_vals = lookup_rows(wb_post, keys)

    print()
    print(f"{'Tutor':<22}{'Date Range':<20}{'PRE-push Tier':<20}{'POST-push Tier':<20}")
    for key in SAMPLE_KEYS:
        pre = pre_vals.get(key, {}).get("Tier", "<row not found>")
        post = post_vals.get(key, {}).get("Tier", "<row not found>")
        print(f"{key[0]:<22}{key[1]:<20}{str(pre):<20}{str(post):<20}")

    print()
    print("Full detail:")
    for key in SAMPLE_KEYS:
        print(f"  {key}:")
        print(f"    PRE-push : {pre_vals.get(key, '<row not found>')}")
        print(f"    POST-push: {post_vals.get(key, '<row not found>')}")


if __name__ == "__main__":
    main()
