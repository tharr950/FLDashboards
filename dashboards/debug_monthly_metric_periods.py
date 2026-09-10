#!/usr/bin/env python3
"""
debug_monthly_metric_periods.py

One-off diagnostic -- NOT part of the regular pipeline. Fetches the
CURRENT (post-push) December_Annual_Reviews.xlsx from GitHub and reports,
for every distinct "Date Range" in MonthlyMetric: row count, and whether
the key KPI columns actually have real (non-blank) values. Used to check
whether appending the 7/12/26-8/8/26 period actually damaged/emptied any
prior period's data, or whether the dashboard app's "previous period"
comparison is failing for some other reason (e.g. how it parses/sorts
Date Range strings) while the underlying data is still intact.

Usage:
    python3 debug_monthly_metric_periods.py
"""

import base64
import io
import os

import openpyxl
import requests
from dotenv import load_dotenv

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(SCRIPT_DIR, ".env"))

GITHUB_TOKEN = os.environ["GITHUB_TOKEN"]
GITHUB_CODE_REPO = os.environ["GITHUB_CODE_REPO"]
HEADERS = {"Authorization": f"token {GITHUB_TOKEN}", "Accept": "application/vnd.github.v3+json"}

FILE_PATH = "December_Annual_Reviews.xlsx"
SHEET_NAME = "MonthlyMetric"


def main():
    url = f"https://api.github.com/repos/{GITHUB_CODE_REPO}/contents/{FILE_PATH}"
    r = requests.get(url, headers=HEADERS, timeout=30)
    r.raise_for_status()
    content = base64.b64decode(r.json()["content"])
    wb = openpyxl.load_workbook(io.BytesIO(content))

    print(f"Sheets in workbook: {wb.sheetnames}")
    ws = wb[SHEET_NAME]
    header = [c.value for c in ws[1]]
    print(f"MonthlyMetric header: {header}")
    print(f"MonthlyMetric total rows (incl header): {ws.max_row}")
    print()

    date_col = header.index("Date Range") + 1
    delivery_col = header.index("% to Delivery Target") + 1
    avail_col = header.index("% to Availability Target") + 1

    from collections import defaultdict
    periods = defaultdict(lambda: {"rows": 0, "blank_delivery": 0, "blank_avail": 0})
    first_row_by_period = {}
    for row in range(2, ws.max_row + 1):
        d = ws.cell(row=row, column=date_col).value
        periods[d]["rows"] += 1
        if ws.cell(row=row, column=delivery_col).value in (None, ""):
            periods[d]["blank_delivery"] += 1
        if ws.cell(row=row, column=avail_col).value in (None, ""):
            periods[d]["blank_avail"] += 1
        if d not in first_row_by_period:
            first_row_by_period[d] = row

    print(f"{'Date Range':<25} {'rows':>6} {'blank delivery':>15} {'blank avail':>12} {'first row #':>12}")
    for d in sorted(periods.keys(), key=lambda x: str(x)):
        p = periods[d]
        print(f"{str(d):<25} {p['rows']:>6} {p['blank_delivery']:>15} {p['blank_avail']:>12} {first_row_by_period[d]:>12}")

    print()
    target = "6/14/26 - 7/11/26"
    if target in periods:
        print(f'"{target}" is present with {periods[target]["rows"]} rows.')
        # Show a couple of actual sample rows for this period.
        shown = 0
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=date_col).value == target:
                print(f'  row {row}: ' + str({h: ws.cell(row=row, column=i+1).value for i, h in enumerate(header)}))
                shown += 1
                if shown >= 2:
                    break
    else:
        print(f'"{target}" is NOT present at all in the sheet -- distinct values were: {sorted(periods.keys(), key=str)}')

    # Also check for other sheets that might drive the dashboard's
    # "Team KPI Changes from Previous Period" table -- if it's reading
    # from somewhere else entirely, that's the real place to look.
    print()
    print("Other sheets (first row of each, for a quick look):")
    for name in wb.sheetnames:
        if name == SHEET_NAME:
            continue
        other = wb[name]
        print(f'  {name}: {other.max_row} rows, header: {[c.value for c in other[1]]}')


if __name__ == "__main__":
    main()
