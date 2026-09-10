#!/usr/bin/env python3
"""
scan_formula_damage.py -- one-off diagnostic, NOT part of the regular
pipeline. Quantifies how many MonthlyMetric rows have Tier /
Adjunct-Professional / Faculty Leader stored as an unevaluated VLOOKUP
formula string (starts with "=") vs a real value, broken down by
period. Also checks whether the embedded MasterTutor sheet in the same
workbook can be used to resolve those broken cells locally (without
needing the external "Master Tutor List" file the formulas point at).

Usage:
    python3 scan_formula_damage.py
"""

import base64
import io
import os
from collections import defaultdict

import openpyxl
import requests
from dotenv import load_dotenv

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(SCRIPT_DIR, ".env"))

GITHUB_TOKEN = os.environ["GITHUB_TOKEN"]
GITHUB_CODE_REPO = os.environ["GITHUB_CODE_REPO"]
HEADERS = {"Authorization": f"token {GITHUB_TOKEN}", "Accept": "application/vnd.github.v3+json"}

FILE_PATH = "December_Annual_Reviews.xlsx"
SHEET_NAME = "MonthlyMetric"


def is_formula(v):
    return isinstance(v, str) and v.startswith("=")


def main():
    url = f"https://api.github.com/repos/{GITHUB_CODE_REPO}/contents/{FILE_PATH}"
    r = requests.get(url, headers=HEADERS, timeout=30)
    r.raise_for_status()
    content = base64.b64decode(r.json()["content"])
    wb = openpyxl.load_workbook(io.BytesIO(content))

    ws = wb[SHEET_NAME]
    header = [c.value for c in ws[1]]
    date_col = header.index("Date Range") + 1
    tutor_col = header.index("Tutor Name") + 1
    tier_col = header.index("Tier") + 1
    adj_col = header.index("Adjunct/Professional") + 1
    fl_col = header.index("Faculty Leader") + 1

    stats = defaultdict(lambda: {
        "rows": 0, "tier_formula": 0, "adj_formula": 0, "fl_formula": 0,
        "tier_blank": 0, "adj_blank": 0, "fl_blank": 0,
    })
    for row in range(2, ws.max_row + 1):
        d = ws.cell(row=row, column=date_col).value
        s = stats[d]
        s["rows"] += 1
        tv = ws.cell(row=row, column=tier_col).value
        av = ws.cell(row=row, column=adj_col).value
        fv = ws.cell(row=row, column=fl_col).value
        if is_formula(tv):
            s["tier_formula"] += 1
        elif tv in (None, ""):
            s["tier_blank"] += 1
        if is_formula(av):
            s["adj_formula"] += 1
        elif av in (None, ""):
            s["adj_blank"] += 1
        if is_formula(fv):
            s["fl_formula"] += 1
        elif fv in (None, ""):
            s["fl_blank"] += 1

    print(f"{'Date Range':<22}{'rows':>6}{'tier=formula':>14}{'tier=blank':>12}{'fl=formula':>12}{'fl=blank':>10}{'adj=formula':>13}{'adj=blank':>11}")
    for d in sorted(stats.keys(), key=str):
        s = stats[d]
        print(f"{str(d):<22}{s['rows']:>6}{s['tier_formula']:>14}{s['tier_blank']:>12}{s['fl_formula']:>12}{s['fl_blank']:>10}{s['adj_formula']:>13}{s['adj_blank']:>11}")

    print()
    if "MasterTutor" not in wb.sheetnames:
        print("No MasterTutor sheet found in this workbook -- can't check local resolution.")
        return

    mt = wb["MasterTutor"]
    mt_header = [c.value for c in mt[1]]
    print(f"MasterTutor header: {mt_header}, rows: {mt.max_row}")
    for row in range(2, min(mt.max_row + 1, 7)):
        print("  ", [mt.cell(row=row, column=i + 1).value for i in range(len(mt_header))])

    mt_lookup = {}
    for row in range(2, mt.max_row + 1):
        name = mt.cell(row=row, column=mt_header.index("Full Name") + 1).value
        if name:
            mt_lookup[name] = {
                "Faculty Leader": mt.cell(row=row, column=mt_header.index("Faculty Leader") + 1).value,
                "Professional/Adjunct": mt.cell(row=row, column=mt_header.index("Professional/Adjunct") + 1).value,
                "Tier": mt.cell(row=row, column=mt_header.index("Tier") + 1).value,
            }

    print()
    print("Sample broken rows resolved against MasterTutor sheet:")
    shown = 0
    unresolved = 0
    total_broken = 0
    for row in range(2, ws.max_row + 1):
        tv = ws.cell(row=row, column=tier_col).value
        if is_formula(tv):
            total_broken += 1
            tutor = ws.cell(row=row, column=tutor_col).value
            resolved = mt_lookup.get(tutor)
            if resolved is None:
                unresolved += 1
            if shown < 8:
                print(f"  row {row} tutor={tutor!r} -> MasterTutor lookup: {resolved}")
                shown += 1
    print()
    print(f"Total rows with broken Tier formula: {total_broken}; of those, {unresolved} have NO match in MasterTutor by exact name.")


if __name__ == "__main__":
    main()
